"""Build the one-slide DA4022 mid-project report.

The slide is designed as a presentation, not a compressed report page:
one claim, four directly comparable results, a compact method diagram,
and a clear statement of completed and remaining work.
"""

import csv
import json
import os
from statistics import mean, pstdev

import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from team import NAMES, ROLLS


HERE = os.path.dirname(os.path.abspath(__file__))
RESULTS = os.path.join(HERE, "results")
PUBLIC = os.environ.get("PENCIL_PUBLIC") == "1"

# Warm editorial palette.
PAPER = RGBColor(0xF7, 0xF3, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x17, 0x1A, 0x1F)
SLATE = RGBColor(0x43, 0x49, 0x52)
MUTED = RGBColor(0x77, 0x7D, 0x85)
RULE = RGBColor(0xD9, 0xD4, 0xCB)
CRIMSON = RGBColor(0x9D, 0x1B, 0x32)
NAVY = RGBColor(0x20, 0x36, 0x57)
TEAL = RGBColor(0x0F, 0x76, 0x70)
AMBER = RGBColor(0xB4, 0x63, 0x18)
PALE_ROSE = RGBColor(0xF4, 0xE7, 0xE9)
PALE_TEAL = RGBColor(0xE4, 0xF0, 0xED)
PALE_AMBER = RGBColor(0xF5, 0xEA, 0xDA)
PALE_BLUE = RGBColor(0xE8, 0xED, 0xF3)


def load_metrics():
    selection = json.load(open(os.path.join(RESULTS, "selection.json"), encoding="utf-8"))
    order = ["cora_pool", "citeseer_pool", "cora_heart", "citeseer_heart"]
    labels = ["Cora\nPool", "Citeseer\nPool", "Cora\nHeaRT", "Citeseer\nHeaRT"]
    metrics = []

    for tag, label in zip(order, labels):
        rows = []
        with open(os.path.join(RESULTS, tag, "ksweep.csv"), newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                if (
                    row["split"] == "test"
                    and row["budget"] == "d2_f10"
                    and abs(float(row["frac"]) - 1.0) < 1e-12
                ):
                    rows.append(float(row["mrr"]) * 100)

        if not rows:
            raise RuntimeError(f"Missing d2_f10 full-candidate rows for {tag}")

        selected_10 = next(r for r in selection[tag]["rows"] if abs(r["cap"] - 0.1) < 1e-12)
        metrics.append(
            {
                "tag": tag,
                "label": label,
                "aa": selection[tag]["aa_test"],
                "frugal": selected_10["test"],
                "frugal_cost": selected_10["cost"] * 100,
                "half": mean(rows),
                "half_std": pstdev(rows),
                "full": selection[tag]["full_test"],
                "full_std": selection[tag]["full_std"],
            }
        )
    return metrics


def make_comparison_figure(metrics):
    out = os.path.join(RESULTS, "fig_mid_v2.png")
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 8,
            "axes.titleweight": "bold",
            "axes.edgecolor": "#D9D4CB",
            "axes.labelcolor": "#434952",
            "xtick.color": "#777D85",
            "ytick.color": "#777D85",
        }
    )
    fig, axes = plt.subplots(1, 4, figsize=(12.4, 2.6), dpi=220)
    fig.patch.set_facecolor("#FFFFFF")
    colors = ["#C9CDD2", "#B46318", "#0F7670", "#203657"]
    names = ["AA", "≤10%", "2×10", "2×20"]

    for ax, item in zip(axes, metrics):
        values = [item["aa"], item["frugal"], item["half"], item["full"]]
        bars = ax.bar(range(4), values, color=colors, width=0.68)
        lower = max(0, min(values) - max(2.0, (max(values) - min(values)) * 0.35))
        upper = max(values) + max(2.0, (max(values) - min(values)) * 0.28)
        ax.set_ylim(lower, upper)
        ax.set_xticks(range(4), names, fontsize=7.3)
        ax.set_title(item["label"].replace("\n", " · "), fontsize=9.2, color="#171A1F", pad=7)
        ax.grid(axis="y", color="#E8E4DD", linewidth=0.65)
        ax.set_axisbelow(True)
        ax.spines[["top", "right", "left"]].set_visible(False)
        ax.tick_params(axis="y", length=0, labelsize=6.8)
        ax.tick_params(axis="x", length=0, pad=4)
        for bar, value in zip(bars, values):
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                value + (upper - lower) * 0.035,
                f"{value:.1f}",
                ha="center",
                va="bottom",
                fontsize=7.3,
                fontweight="bold",
                color="#171A1F",
            )

    fig.text(0.009, 0.54, "Test MRR ×100", rotation=90, va="center", color="#777D85", fontsize=7.5)
    fig.text(
        0.5,
        0.01,
        "AA = CPU heuristic  ·  ≤10% = validation-selected hybrid  ·  2×10 = half fanout  ·  2×20 = official full PENCIL",
        ha="center",
        color="#5E646D",
        fontsize=7.2,
    )
    fig.subplots_adjust(left=0.045, right=0.995, top=0.82, bottom=0.25, wspace=0.28)
    fig.savefig(out, facecolor="white", bbox_inches="tight", pad_inches=0.04)
    plt.close(fig)
    return out


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])


def shape(kind, x, y, w, h, fill=None, line=None, line_width=0.7, radius=False):
    if radius:
        kind = MSO_SHAPE.ROUNDED_RECTANGLE
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_width)
    if radius:
        try:
            sh.adjustments[0] = 0.08
        except Exception:
            pass
    return sh


def textbox(x, y, w, h, text="", margin=0.0, valign=MSO_ANCHOR.TOP):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    if text:
        tf.paragraphs[0].text = text
    return sh


def paragraph(
    tf,
    runs,
    *,
    first=False,
    align=PP_ALIGN.LEFT,
    line_spacing=1.0,
    space_after=0,
    space_before=0,
):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line_spacing
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    for text, size, bold, color, font, italic in runs:
        r = p.add_run()
        r.text = text
        r.font.name = font
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = color
    return p


def run(text, size, bold=False, color=INK, font="Aptos", italic=False):
    return (text, size, bold, color, font, italic)


def rule(x, y, w, color=RULE, height=0.012):
    shape(MSO_SHAPE.RECTANGLE, x, y, w, height, fill=color)


metrics = load_metrics()
chart_path = make_comparison_figure(metrics)

# Canvas and structural accents.
shape(MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 7.5, fill=PAPER)
shape(MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.075, fill=CRIMSON)
shape(MSO_SHAPE.RECTANGLE, 0.27, 0.31, 0.055, 6.87, fill=CRIMSON)

# Header.
tf = textbox(0.55, 0.25, 8.8, 0.24).text_frame
paragraph(
    tf,
    [run("DA4022  /  MID-PROJECT REPORT  /  15 OCT 2026", 8.5, True, CRIMSON, "Aptos")],
    first=True,
)

tf = textbox(0.55, 0.50, 8.35, 0.43).text_frame
paragraph(
    tf,
    [run("Full-budget PENCIL is", 28, True, INK, "Georgia")],
    first=True,
)
tf = textbox(0.55, 1.00, 8.35, 0.43).text_frame
paragraph(
    tf,
    [run("often unnecessary.", 28, True, CRIMSON, "Georgia")],
    first=True,
)

tf = textbox(9.05, 0.27, 3.72, 0.95).text_frame
paragraph(
    tf,
    [run("HOW MUCH PENCIL CAN YOU AFFORD?", 8.3, True, MUTED, "Aptos")],
    first=True,
    align=PP_ALIGN.RIGHT,
)
paragraph(
    tf,
    [run("Compute-allocation laws for link prediction", 12.2, True, NAVY, "Georgia")],
    align=PP_ALIGN.RIGHT,
    space_before=3,
    line_spacing=1.02,
)

rule(0.55, 1.49, 12.23)

# Hero statement.
shape(MSO_SHAPE.ROUNDED_RECTANGLE, 0.55, 1.72, 12.23, 1.15, fill=NAVY, radius=True)
tf = textbox(0.82, 1.83, 2.15, 0.86, valign=MSO_ANCHOR.MIDDLE).text_frame
paragraph(
    tf,
    [run("65%", 42, True, WHITE, "Georgia")],
    first=True,
    align=PP_ALIGN.CENTER,
)
tf = textbox(2.97, 1.91, 0.95, 0.62, valign=MSO_ANCHOR.MIDDLE).text_frame
paragraph(
    tf,
    [run("GPU\nbudget", 11.2, True, PALE_BLUE, "Aptos")],
    first=True,
    line_spacing=1.05,
)
shape(MSO_SHAPE.RECTANGLE, 4.05, 1.96, 0.015, 0.65, fill=RGBColor(0x5C, 0x6C, 0x83))
tf = textbox(4.32, 1.86, 5.62, 0.72, valign=MSO_ANCHOR.MIDDLE).text_frame
paragraph(
    tf,
    [run("matches or beats full PENCIL", 20, True, WHITE, "Georgia")],
    first=True,
)
paragraph(
    tf,
    [run("in 3 of 4 test regimes — with ~40% less peak VRAM on Cora HeaRT.", 10.3, False, PALE_BLUE, "Aptos")],
    space_before=3,
)
shape(MSO_SHAPE.ROUNDED_RECTANGLE, 10.34, 1.96, 2.15, 0.63, fill=RGBColor(0x2B, 0x4A, 0x70), radius=True)
tf = textbox(10.47, 2.04, 1.89, 0.44, valign=MSO_ANCHOR.MIDDLE).text_frame
paragraph(
    tf,
    [run("s = 2×10  vs  2×20", 10.2, True, WHITE, "Consolas")],
    first=True,
    align=PP_ALIGN.CENTER,
)

# Left column: question, method, progress.
LX, LW = 0.55, 3.48
tf = textbox(LX, 3.12, LW, 0.20).text_frame
paragraph(tf, [run("THE QUESTION", 8.5, True, CRIMSON, "Aptos")], first=True)
tf = textbox(LX, 3.36, LW, 0.72).text_frame
paragraph(
    tf,
    [run("Can we spend PENCIL's GPU compute only where a cheap graph heuristic is uncertain?", 13.2, True, INK, "Georgia")],
    first=True,
    line_spacing=1.03,
)

tf = textbox(LX, 4.20, LW, 0.20).text_frame
paragraph(tf, [run("OUR INFERENCE-ONLY METHOD", 8.5, True, CRIMSON, "Aptos")], first=True)

steps = [
    ("1", "Rank", "Adamic–Adar\non CPU", PALE_AMBER, AMBER),
    ("2", "Shortlist", "top k candidate\npairs", PALE_ROSE, CRIMSON),
    ("3", "Re-rank", "PENCIL with\nbudget s", PALE_BLUE, NAVY),
]
for i, (num, title, body, bg, accent) in enumerate(steps):
    x = LX + i * 1.15
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, 4.48, 1.04, 0.94, fill=bg, line=RULE, radius=True)
    shape(MSO_SHAPE.OVAL, x + 0.08, 4.58, 0.22, 0.22, fill=accent)
    tf = textbox(x + 0.08, 4.575, 0.22, 0.20, valign=MSO_ANCHOR.MIDDLE).text_frame
    paragraph(tf, [run(num, 7.2, True, WHITE, "Aptos")], first=True, align=PP_ALIGN.CENTER)
    tf = textbox(x + 0.35, 4.55, 0.60, 0.22).text_frame
    paragraph(tf, [run(title, 8.5, True, accent, "Aptos")], first=True)
    tf = textbox(x + 0.08, 4.86, 0.88, 0.42).text_frame
    paragraph(tf, [run(body, 7.5, False, SLATE, "Aptos")], first=True, align=PP_ALIGN.CENTER, line_spacing=0.95)
    if i < 2:
        tf = textbox(x + 1.04, 4.78, 0.11, 0.22).text_frame
        paragraph(tf, [run("›", 16, True, MUTED, "Georgia")], first=True, align=PP_ALIGN.CENTER)

tf = textbox(LX, 5.66, LW, 0.20).text_frame
paragraph(tf, [run("MID-PROJECT STATUS", 8.5, True, CRIMSON, "Aptos")], first=True)
status_lines = [
    ("✓", "Reproduced PENCIL", "Cora 44.00 · Citeseer 50.46"),
    ("✓", "112 evaluation runs", "2 graphs × 2 protocols × budgets"),
    ("→", "Next: final decision rule", "latency/VRAM Pareto + final report"),
]
for i, (mark, title, detail) in enumerate(status_lines):
    y = 5.93 + i * 0.37
    color = TEAL if mark == "✓" else AMBER
    tf = textbox(LX, y, 0.25, 0.22).text_frame
    paragraph(tf, [run(mark, 10.5, True, color, "Aptos")], first=True)
    tf = textbox(LX + 0.29, y, 3.15, 0.25).text_frame
    paragraph(
        tf,
        [run(title, 8.5, True, INK, "Aptos"), run("  " + detail, 7.5, False, MUTED, "Aptos")],
        first=True,
    )

# Right column: four result cards.
RX, RW = 4.32, 8.46
tf = textbox(RX, 3.12, RW, 0.20).text_frame
paragraph(
    tf,
    [run("HALF-FANOUT RESULT  /  TEST MRR ×100", 8.5, True, CRIMSON, "Aptos")],
    first=True,
)

card_gap = 0.12
card_w = (RW - 3 * card_gap) / 4
for i, item in enumerate(metrics):
    x = RX + i * (card_w + card_gap)
    delta = item["half"] - item["full"]
    if delta > 0.25:
        verdict, bg, accent = "BEATS FULL", PALE_TEAL, TEAL
    elif delta > -1.0:
        verdict, bg, accent = "MATCHES", PALE_BLUE, NAVY
    else:
        verdict, bg, accent = "NEEDS FULL", PALE_ROSE, CRIMSON

    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, 3.40, card_w, 1.42, fill=WHITE, line=RULE, radius=True)
    shape(MSO_SHAPE.RECTANGLE, x, 3.40, card_w, 0.055, fill=accent)
    tf = textbox(x + 0.12, 3.54, card_w - 0.24, 0.24).text_frame
    paragraph(
        tf,
        [run(item["label"].replace("\n", " · "), 8.7, True, SLATE, "Aptos")],
        first=True,
        align=PP_ALIGN.CENTER,
    )
    tf = textbox(x + 0.10, 3.83, card_w - 0.20, 0.42).text_frame
    paragraph(
        tf,
        [run(f"{item['half']:.1f}", 22, True, accent, "Georgia")],
        first=True,
        align=PP_ALIGN.CENTER,
    )
    tf = textbox(x + 0.10, 4.23, card_w - 0.20, 0.20).text_frame
    paragraph(
        tf,
        [run(f"full: {item['full']:.1f}", 7.6, False, MUTED, "Consolas")],
        first=True,
        align=PP_ALIGN.CENTER,
    )
    shape(MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.32, 4.49, card_w - 0.64, 0.23, fill=bg, radius=True)
    tf = textbox(x + 0.34, 4.515, card_w - 0.68, 0.17, valign=MSO_ANCHOR.MIDDLE).text_frame
    paragraph(tf, [run(verdict, 7.0, True, accent, "Aptos")], first=True, align=PP_ALIGN.CENTER)

# Evidence chart.
tf = textbox(RX, 5.03, RW, 0.20).text_frame
paragraph(
    tf,
    [run("THE FULL BUDGET IS NOT THE DEFAULT WINNER", 8.5, True, CRIMSON, "Aptos")],
    first=True,
)
shape(MSO_SHAPE.ROUNDED_RECTANGLE, RX, 5.30, RW, 1.47, fill=WHITE, line=RULE, radius=True)
slide.shapes.add_picture(chart_path, Inches(RX + 0.09), Inches(5.36), width=Inches(RW - 0.18), height=Inches(1.34))

# Bottom takeaway and authors.
rule(0.55, 7.04, 12.23)
tf = textbox(0.55, 7.13, 8.3, 0.20).text_frame
paragraph(
    tf,
    [
        run("INTERIM LAW: ", 8.2, True, CRIMSON, "Aptos"),
        run("reduce neighborhood fanout first; reserve full context for the hard regime.", 8.2, True, INK, "Aptos"),
    ],
    first=True,
)

team_parts = []
for name in NAMES:
    text = name
    if not PUBLIC:
        text += f" ({ROLLS[name]})"
    team_parts.append(text)
tf = textbox(8.60, 7.10, 4.18, 0.24).text_frame
paragraph(
    tf,
    [run(" · ".join(team_parts), 6.4, False, MUTED, "Aptos")],
    first=True,
    align=PP_ALIGN.RIGHT,
)

out = os.path.join(HERE, "docs", "MIDREPORT_public.pptx") if PUBLIC else os.path.join(HERE, "MIDREPORT.pptx")
prs.save(out)
print(f"wrote {out}")
